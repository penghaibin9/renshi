"""Production regression tests for the durable HR05 Excel staging ledger."""

from __future__ import annotations

import io
from unittest.mock import patch

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from hr_onboarding.models import HrOnboardingImportJob, HrOnboardingImportRow
from hr_onboarding.services.excel_service import (
    ExcelImportStateError,
    commit_persisted_import,
    confirm_persisted_import,
    get_persisted_import,
    stage_persisted_import,
)


MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_upload(*, source_type="LEGACY_MIGRATION", source_id="", name="import.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([
        "legal_name", "source_type", "source_id", "expected_report_date",
        "employment_type", "staff_category", "note",
    ])
    ws.append(["张三", source_type, source_id, "2026-09-01", "FULL_TIME", "TEACHER", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue(), content_type=MIME)


class DurableExcelImportTests(TestCase):
    def test_stage_persists_job_rows_and_source_digest(self):
        job = stage_persisted_import(
            tenant_id=1001,
            uploaded_by=7,
            uploaded_file=workbook_upload(),
        )
        self.assertEqual(job.status, HrOnboardingImportJob.Status.READY_TO_COMMIT)
        self.assertEqual(job.row_count, 1)
        self.assertEqual(len(job.source_sha256), 64)
        row = HrOnboardingImportRow.objects.get(job=job)
        self.assertEqual(row.tenant_id, 1001)
        self.assertEqual(row.status, HrOnboardingImportRow.Status.VALID)
        self.assertNotIn("_row", row.payload_json)

    def test_non_legacy_source_without_id_is_durable_validation_failure(self):
        job = stage_persisted_import(
            tenant_id=1001,
            uploaded_by=7,
            uploaded_file=workbook_upload(source_type="HR04_HIRE", source_id=""),
        )
        self.assertEqual(job.status, HrOnboardingImportJob.Status.VALIDATION_FAILED)
        row = HrOnboardingImportRow.objects.get(job=job)
        self.assertEqual(row.status, HrOnboardingImportRow.Status.INVALID)
        self.assertTrue(any(item.get("field") == "source_id" for item in row.error_json))

    def test_job_lookup_is_tenant_scoped(self):
        job = stage_persisted_import(
            tenant_id=1001,
            uploaded_by=7,
            uploaded_file=workbook_upload(),
        )
        self.assertIsNotNone(get_persisted_import(tenant_id=1001, job_id=job.id))
        self.assertIsNone(get_persisted_import(tenant_id=2002, job_id=job.id))


    def test_confirm_only_queues_job_for_background_worker(self):
        job = stage_persisted_import(
            tenant_id=1001, uploaded_by=7, uploaded_file=workbook_upload()
        )
        queued = confirm_persisted_import(
            tenant_id=1001, job_id=job.id, actor_user_id=7
        )
        self.assertEqual(queued.status, HrOnboardingImportJob.Status.COMMITTING)
        self.assertIsNone(queued.commit_started_at)
        self.assertEqual(queued.confirmed_by, 7)
        self.assertIsNotNone(queued.confirmed_at)
        replay = confirm_persisted_import(
            tenant_id=1001, job_id=job.id, actor_user_id=7
        )
        self.assertEqual(replay.status, HrOnboardingImportJob.Status.COMMITTING)
        self.assertIsNone(replay.commit_started_at)

    @patch("hr_onboarding.services.case_service.CaseService.create_case_from_handoff")
    def test_commit_uses_stable_row_idempotency_key_and_survives_replay(self, create_case):
        create_case.return_value = {"created": True, "case_id": "CASE-1"}
        job = stage_persisted_import(
            tenant_id=1001,
            uploaded_by=7,
            uploaded_file=workbook_upload(),
        )
        committed = commit_persisted_import(
            tenant_id=1001,
            job_id=job.id,
            actor_user_id=7,
        )
        self.assertEqual(committed.status, HrOnboardingImportJob.Status.COMPLETED)
        row = HrOnboardingImportRow.objects.get(job=job)
        self.assertEqual(row.status, HrOnboardingImportRow.Status.CREATED)
        request = create_case.call_args.args[0]
        self.assertEqual(request["source_id"], f"excel:{job.id}:{row.row_no}")
        self.assertEqual(
            create_case.call_args.kwargs["idempotency_key"],
            f"hr05-excel:{job.id}:{row.row_no}",
        )
        replay = commit_persisted_import(tenant_id=1001, job_id=job.id, actor_user_id=7)
        self.assertEqual(replay.status, HrOnboardingImportJob.Status.COMPLETED)
        self.assertEqual(create_case.call_count, 1)


    def test_missing_required_header_is_rejected_before_job_creation(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["legal_name", "source_type"])
        ws.append(["张三", "LEGACY_MIGRATION"])
        buf = io.BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile("missing.xlsx", buf.getvalue(), content_type=MIME)
        with self.assertRaises(ExcelImportStateError) as ctx:
            stage_persisted_import(tenant_id=1001, uploaded_by=7, uploaded_file=upload)
        self.assertEqual(ctx.exception.code, "EXCEL_REQUIRED_HEADER_MISSING")
        self.assertFalse(HrOnboardingImportJob.objects.exists())

    def test_duplicate_header_is_rejected_before_job_creation(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["legal_name", "source_type", "expected_report_date", "legal_name"])
        ws.append(["张三", "LEGACY_MIGRATION", "2026-09-01", "重复"])
        buf = io.BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile("duplicate.xlsx", buf.getvalue(), content_type=MIME)
        with self.assertRaises(ExcelImportStateError) as ctx:
            stage_persisted_import(tenant_id=1001, uploaded_by=7, uploaded_file=upload)
        self.assertEqual(ctx.exception.code, "EXCEL_DUPLICATE_HEADER")

    def test_header_only_workbook_is_rejected(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([
            "legal_name", "source_type", "source_id", "expected_report_date",
            "employment_type", "staff_category", "note",
        ])
        buf = io.BytesIO()
        wb.save(buf)
        upload = SimpleUploadedFile("empty.xlsx", buf.getvalue(), content_type=MIME)
        with self.assertRaises(ExcelImportStateError) as ctx:
            stage_persisted_import(tenant_id=1001, uploaded_by=7, uploaded_file=upload)
        self.assertEqual(ctx.exception.code, "EXCEL_EMPTY_WORKBOOK")

    def test_rejects_non_xlsx_before_staging(self):
        upload = SimpleUploadedFile("bad.csv", b"a,b\n1,2", content_type="text/csv")
        with self.assertRaises(ExcelImportStateError) as ctx:
            stage_persisted_import(tenant_id=1001, uploaded_by=7, uploaded_file=upload)
        self.assertEqual(ctx.exception.code, "EXCEL_FILE_TYPE_INVALID")
