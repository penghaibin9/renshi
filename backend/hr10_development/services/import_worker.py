"""Controlled, lease-safe and idempotent Excel/legacy import pipeline for HR10."""

from __future__ import annotations

import hashlib
import json
import logging
import uuid
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

logger = logging.getLogger(__name__)

MAX_IMPORT_ROWS = 10_000
SUPPORTED_TEMPLATE_VERSION = "V1"
DEFAULT_LEASE_SECONDS = 300
HEARTBEAT_EVERY_ROWS = 250

TEMPLATE_SCHEMAS = {
    "EXCEL_PLAN": {
        "required": ("plan_no", "plan_type", "start_date", "end_date"),
        "date_fields": ("start_date", "end_date"),
        "integer_fields": ("owner_org_id",),
        "business_key": "plan_no",
        "target_model": "HrDevelopmentPlan",
    },
    "EXCEL_PROGRAM": {
        "required": ("program_code", "title", "activity_type"),
        "date_fields": (),
        "integer_fields": ("owner_org_id", "provider_org_id"),
        "business_key": "program_code",
        "target_model": "HrLearningProgram",
    },
    "EXCEL_PRACTICE": {
        "required": ("project_no", "title", "provider_org_id"),
        "date_fields": ("planned_start_date", "planned_end_date"),
        "integer_fields": ("provider_org_id", "owner_org_id", "capacity"),
        "business_key": "project_no",
        "target_model": "HrEnterprisePracticeProject",
    },
}

HEADER_ALIASES = {
    "计划编号": "plan_no",
    "计划类型": "plan_type",
    "开始日期": "start_date",
    "结束日期": "end_date",
    "归属组织ID": "owner_org_id",
    "教职工ID": "staff_master_id",
    "项目编码": "program_code",
    "项目编号": "project_no",
    "项目标题": "title",
    "活动类型": "activity_type",
    "机构ID": "provider_org_id",
    "计划开始日期": "planned_start_date",
    "计划结束日期": "planned_end_date",
    "容量": "capacity",
}

_PARSE_ACTIVE = {"PARSE", "VALIDATION"}
_EXECUTE_ACTIVE = {"CONFIRMING", "EXECUTING"}
_TERMINAL = {"SUCCESS", "FAILED", "CANCELLED"}


class ImportLeaseLost(RuntimeError):
    """The current worker no longer owns the job claim."""


class ImportJobBusy(RuntimeError):
    """Another worker currently owns an unexpired import lease."""

    def __init__(self, job):
        super().__init__(f"IMPORT_JOB_BUSY:{job.status}")
        self.job = job


class ImportExecutionError(RuntimeError):
    def __init__(self, code: str, message: str, *, row_id: int | None = None, row_number: int | None = None):
        super().__init__(message)
        self.code = code
        self.row_id = row_id
        self.row_number = row_number


def _lease_seconds() -> int:
    configured = int(getattr(settings, "HR10_IMPORT_LEASE_SECONDS", DEFAULT_LEASE_SECONDS) or DEFAULT_LEASE_SECONDS)
    return min(max(configured, 30), 3600)


def _lease_deadline(now=None):
    now = now or timezone.now()
    return now + timedelta(seconds=_lease_seconds())


def _lease_is_active(job, now=None) -> bool:
    now = now or timezone.now()
    return bool(job.claim_token and job.lease_expires_at and job.lease_expires_at > now)


def _claim_parse_job(job_id: int):
    from hr10_development.legacy.import_job import HrDevelopmentImportJob

    now = timezone.now()
    token = uuid.uuid4().hex
    with transaction.atomic():
        job = HrDevelopmentImportJob.objects.select_for_update().get(id=job_id)
        if job.status in {"PREVIEW", *_TERMINAL}:
            return job, None
        if job.status in _EXECUTE_ACTIVE:
            if _lease_is_active(job, now):
                return job, None
            return job, None
        if job.status not in {"PENDING", *_PARSE_ACTIVE}:
            return job, None
        if job.status in _PARSE_ACTIVE and _lease_is_active(job, now):
            return job, None

        takeover = job.status in _PARSE_ACTIVE
        summary = dict(job.result_summary_json or {})
        if takeover:
            summary["leaseTakeovers"] = int(summary.get("leaseTakeovers", 0)) + 1

        job.status = "PARSE"
        job.claim_token = token
        job.heartbeat_at = now
        job.lease_expires_at = _lease_deadline(now)
        job.started_at = job.started_at or now
        job.completed_at = None
        job.retry_count += 1
        job.result_summary_json = summary
        job.save(
            update_fields=[
                "status",
                "claim_token",
                "heartbeat_at",
                "lease_expires_at",
                "started_at",
                "completed_at",
                "retry_count",
                "result_summary_json",
                "updated_at",
            ]
        )
        return job, token


def _heartbeat(job_id: int, claim_token: str, *, statuses=None) -> None:
    from hr10_development.legacy.import_job import HrDevelopmentImportJob

    allowed = set(statuses or _PARSE_ACTIVE)
    now = timezone.now()
    updated = HrDevelopmentImportJob.objects.filter(
        id=job_id,
        claim_token=claim_token,
        status__in=allowed,
    ).update(heartbeat_at=now, lease_expires_at=_lease_deadline(now))
    if updated != 1:
        raise ImportLeaseLost("IMPORT_LEASE_LOST")


def _require_claim_locked(job, claim_token: str, *, statuses) -> None:
    if job.claim_token != claim_token or job.status not in set(statuses):
        raise ImportLeaseLost("IMPORT_LEASE_LOST")


def _finish_parse(job_id: int, claim_token: str, final_status: str):
    from hr10_development.legacy.import_job import HrDevelopmentImportJob

    with transaction.atomic():
        job = HrDevelopmentImportJob.objects.select_for_update().get(id=job_id)
        _require_claim_locked(job, claim_token, statuses=_PARSE_ACTIVE)
        job.status = final_status
        job.claim_token = ""
        job.lease_expires_at = None
        job.completed_at = timezone.now() if final_status == "SUCCESS" else None
        job.save(
            update_fields=[
                "status",
                "claim_token",
                "lease_expires_at",
                "completed_at",
                "updated_at",
            ]
        )
        return job


def _fail_claimed_job(job_id: int, claim_token: str, exc: Exception, *, phase: str):
    from hr10_development.legacy.import_job import HrDevelopmentImportJob

    now = timezone.now()
    with transaction.atomic():
        job = HrDevelopmentImportJob.objects.select_for_update().filter(id=job_id).first()
        if not job or job.claim_token != claim_token:
            return job
        summary = dict(job.result_summary_json or {})
        summary.update(
            {
                "errorCode": "IMPORT_PARSE_FAILED" if phase == "PARSE" else "IMPORT_EXECUTION_FAILED",
                "error": str(exc)[:2000],
                "failedPhase": phase,
            }
        )
        job.status = "FAILED"
        job.result_summary_json = summary
        job.completed_at = now
        job.claim_token = ""
        job.lease_expires_at = None
        job.save(
            update_fields=[
                "status",
                "result_summary_json",
                "completed_at",
                "claim_token",
                "lease_expires_at",
                "updated_at",
            ]
        )
        return job


def run_import_job(job_id: int):
    """Parse one job only after obtaining a renewable lease claim.

    An active lease makes duplicate worker delivery a no-op. If a worker dies,
    the same job becomes reclaimable only after the lease expires. Completion
    is accepted only from the worker that still owns the claim token.
    """

    from hr10_development.legacy.import_job import HrDevelopmentImportJob

    claim_token = None
    try:
        job, claim_token = _claim_parse_job(job_id)
        if not claim_token:
            return job

        if "LEGACY_EMPLOYEE" in job.job_type:
            _parse_legacy_employee(job, claim_token)
            final_status = "SUCCESS"
        elif job.job_type in TEMPLATE_SCHEMAS:
            _parse_excel(job, claim_token)
            final_status = "PREVIEW"
        else:
            raise ValueError(f"UNSUPPORTED_JOB_TYPE: {job.job_type}")

        return _finish_parse(job.id, claim_token, final_status)
    except HrDevelopmentImportJob.DoesNotExist:
        logger.error("Import job %s not found", job_id)
        return None
    except ImportLeaseLost:
        logger.warning("Import job %s lost lease; stale worker result discarded", job_id)
        return HrDevelopmentImportJob.objects.filter(id=job_id).first()
    except Exception as exc:
        logger.exception("Import job %s failed", job_id)
        if claim_token:
            _fail_claimed_job(job_id, claim_token, exc, phase="PARSE")
        return HrDevelopmentImportJob.objects.filter(id=job_id).first()


def _parse_legacy_employee(job, claim_token: str):
    """Stage legacy Employee qualification values without silently claiming zero."""

    from employee.models import Employee
    from hr10_development.legacy.import_job import HrDevelopmentImportJob
    from hr10_development.legacy.staging import HrDevelopmentStagingRow

    employees = Employee.objects.filter(tenant_id=job.tenant_id, is_active=True)[:5000]
    created = 0
    for index, emp in enumerate(employees, start=1):
        if index % HEARTBEAT_EVERY_ROWS == 0:
            _heartbeat(job.id, claim_token)
        if not emp.qualification:
            continue
        HrDevelopmentStagingRow.objects.update_or_create(
            tenant_id=job.tenant_id,
            import_job_id=job.id,
            source_object_id=str(emp.id),
            defaults={
                "source_system": "LEGACY_EMPLOYEE",
                "source_table": "Employee",
                "source_field": "qualification",
                "raw_text": emp.qualification,
                "migration_trust_level": "UNKNOWN",
                "verification_status": "PENDING",
                "execution_status": "PENDING",
            },
        )
        created += 1

    with transaction.atomic():
        locked = HrDevelopmentImportJob.objects.select_for_update().get(id=job.id)
        _require_claim_locked(locked, claim_token, statuses=_PARSE_ACTIVE)
        locked.total_rows = created
        locked.processed_rows = created
        locked.error_rows = 0
        locked.result_summary_json = {
            **dict(locked.result_summary_json or {}),
            "stagedRows": created,
            "sourceStatus": "AVAILABLE",
            "message": "Legacy Employee data staged for review",
        }
        locked.heartbeat_at = timezone.now()
        locked.lease_expires_at = _lease_deadline(locked.heartbeat_at)
        locked.save(
            update_fields=[
                "total_rows",
                "processed_rows",
                "error_rows",
                "result_summary_json",
                "heartbeat_at",
                "lease_expires_at",
                "updated_at",
            ]
        )


def _normalise_header(value) -> str:
    header = str(value or "").strip()
    return HEADER_ALIASES.get(header, header.lower().replace(" ", "_"))


def _json_value(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _coerce_row(raw: dict, schema: dict) -> tuple[dict, list[str]]:
    parsed = {key: _json_value(value) for key, value in raw.items() if value not in (None, "")}
    errors = []

    for field in schema["required"]:
        if parsed.get(field) in (None, ""):
            errors.append(f"{field}: REQUIRED")

    for field in schema["date_fields"]:
        value = parsed.get(field)
        if value in (None, ""):
            continue
        try:
            parsed[field] = date.fromisoformat(str(value)[:10]).isoformat()
        except ValueError:
            errors.append(f"{field}: INVALID_DATE")

    for field in schema["integer_fields"]:
        value = parsed.get(field)
        if value in (None, ""):
            continue
        try:
            parsed[field] = int(value)
        except (TypeError, ValueError):
            errors.append(f"{field}: INVALID_INTEGER")

    start = parsed.get("start_date") or parsed.get("planned_start_date")
    end = parsed.get("end_date") or parsed.get("planned_end_date")
    if start and end and start > end:
        errors.append("date_range: START_AFTER_END")
    if isinstance(parsed.get("capacity"), int) and parsed["capacity"] < 0:
        errors.append("capacity: NEGATIVE")
    return parsed, errors


def _batch_authority_validation_errors(job_type: str, tenant_id: int, rows: list[tuple[int, dict, dict]]) -> dict[int, list[str]]:
    """Batch authority checks so a 10k-row preview does not become an N+1 query storm."""

    from hr10_development.constants import DevelopmentActivityType, PlanType
    from hr10_development.models import (
        HrDevelopmentPlan,
        HrDevelopmentProviderOrganization,
        HrEnterprisePracticeProject,
        HrLearningProgram,
    )
    from hr_structure.models import HrOrganization

    errors_by_row: dict[int, list[str]] = {row_number: [] for row_number, _, _ in rows}
    schema = TEMPLATE_SCHEMAS[job_type]
    business_key = schema["business_key"]
    target_lookup = {
        "EXCEL_PLAN": (HrDevelopmentPlan, "plan_no"),
        "EXCEL_PROGRAM": (HrLearningProgram, "program_code"),
        "EXCEL_PRACTICE": (HrEnterprisePracticeProject, "project_no"),
    }
    model, field = target_lookup[job_type]
    requested_keys = {str(parsed.get(business_key, "") or "").strip() for _, _, parsed in rows}
    requested_keys.discard("")
    existing_keys = set(
        model.objects.filter(tenant_id=tenant_id, **{f"{field}__in": requested_keys}).values_list(field, flat=True)
    ) if requested_keys else set()

    owner_org_ids = {parsed.get("owner_org_id") for _, _, parsed in rows if parsed.get("owner_org_id") is not None}
    valid_owner_org_ids = set(
        HrOrganization.objects.filter(tenant_id=tenant_id, id__in=owner_org_ids).values_list("id", flat=True)
    ) if owner_org_ids else set()

    provider_org_ids = {parsed.get("provider_org_id") for _, _, parsed in rows if parsed.get("provider_org_id") is not None}
    valid_provider_org_ids = set(
        HrDevelopmentProviderOrganization.objects.filter(
            tenant_id=tenant_id, id__in=provider_org_ids
        ).values_list("id", flat=True)
    ) if provider_org_ids else set()

    # Resolve individual-plan staff references in one tenant-scoped batch.
    # The Excel column keeps its historic name for template compatibility, but
    # values may now be canonical HR03 UUIDs or legacy Employee ids.
    staff_refs = {
        str(parsed.get("staff_master_id") or "").strip()
        for _, _, parsed in rows
        if job_type == "EXCEL_PLAN" and parsed.get("plan_type") == PlanType.INDIVIDUAL
        and parsed.get("staff_master_id") not in (None, "")
    }
    uuid_refs = set()
    legacy_refs = set()
    invalid_staff_refs = set()
    for raw_ref in staff_refs:
        try:
            uuid_refs.add(uuid.UUID(raw_ref))
            continue
        except (ValueError, TypeError, AttributeError):
            pass
        try:
            legacy = int(raw_ref)
            if legacy <= 0:
                raise ValueError
            legacy_refs.add(legacy)
        except (ValueError, TypeError):
            invalid_staff_refs.add(raw_ref)

    staff_by_uuid = {}
    staff_by_legacy = {}
    ambiguous_legacy = set()
    if uuid_refs or legacy_refs:
        from django.db.models import Q
        from hr_staff.models import HrStaffMaster
        staff_rows = HrStaffMaster.objects.filter(tenant_id=tenant_id).filter(
            Q(id__in=uuid_refs) | Q(legacy_employee_id__in=legacy_refs)
        ).only("id", "legacy_employee_id")
        for staff in staff_rows:
            staff_by_uuid[str(staff.id)] = staff
            if staff.legacy_employee_id is not None:
                legacy = int(staff.legacy_employee_id)
                if legacy in staff_by_legacy and staff_by_legacy[legacy].id != staff.id:
                    ambiguous_legacy.add(legacy)
                staff_by_legacy[legacy] = staff

    def normalize_staff_ref(parsed):
        raw_ref = str(parsed.get("staff_master_id") or "").strip()
        if raw_ref in invalid_staff_refs:
            return "INVALID"
        staff = staff_by_uuid.get(raw_ref)
        if staff is None:
            try:
                legacy = int(raw_ref)
            except (TypeError, ValueError):
                return "NOT_FOUND"
            if legacy in ambiguous_legacy:
                return "AMBIGUOUS"
            staff = staff_by_legacy.get(legacy)
        if staff is None:
            return "NOT_FOUND"
        parsed["staff_master_uuid"] = str(staff.id)
        parsed["staff_master_legacy_id"] = staff.legacy_employee_id
        return "OK"

    for row_number, _, parsed in rows:
        row_errors = errors_by_row[row_number]
        key_value = str(parsed.get(business_key, "") or "").strip()
        if key_value in existing_keys:
            row_errors.append(f"{business_key}: ALREADY_EXISTS")
        owner_org_id = parsed.get("owner_org_id")
        if owner_org_id is not None and owner_org_id not in valid_owner_org_ids:
            row_errors.append("owner_org_id: NOT_FOUND_IN_TENANT")
        provider_org_id = parsed.get("provider_org_id")
        if provider_org_id is not None and provider_org_id not in valid_provider_org_ids:
            row_errors.append("provider_org_id: NOT_FOUND_IN_TENANT")
        if job_type == "EXCEL_PLAN":
            if parsed.get("plan_type") not in set(PlanType.values):
                row_errors.append("plan_type: INVALID_CHOICE")
            elif parsed.get("plan_type") == PlanType.INDIVIDUAL:
                if not parsed.get("staff_master_id"):
                    row_errors.append("staff_master_id: REQUIRED_FOR_INDIVIDUAL")
                else:
                    identity_status = normalize_staff_ref(parsed)
                    if identity_status == "INVALID":
                        row_errors.append("staff_master_id: INVALID_IDENTITY")
                    elif identity_status == "AMBIGUOUS":
                        row_errors.append("staff_master_id: AMBIGUOUS_LEGACY_MAPPING")
                    elif identity_status == "NOT_FOUND":
                        row_errors.append("staff_master_id: NOT_FOUND_IN_TENANT")
        elif job_type == "EXCEL_PROGRAM" and parsed.get("activity_type") not in set(DevelopmentActivityType.values):
            row_errors.append("activity_type: INVALID_CHOICE")

    return errors_by_row


def _sha256(field_file) -> str:
    digest = hashlib.sha256()
    field_file.open("rb")
    try:
        for chunk in iter(lambda: field_file.read(1024 * 1024), b""):
            digest.update(chunk)
    finally:
        field_file.close()
    return digest.hexdigest()


def _save_error_workbook(job, headers: list[str], errors: list[dict]) -> str:
    if not errors:
        if job.error_workbook_path and default_storage.exists(job.error_workbook_path):
            default_storage.delete(job.error_workbook_path)
        return ""

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("errors")
    sheet.append(["row_number", *headers, "errors"])
    for item in errors:
        sheet.append(
            [item["rowNumber"], *[item["raw"].get(header) for header in headers], "; ".join(item["errors"])]
        )
    payload = BytesIO()
    workbook.save(payload)
    path = f"hr10/imports/{job.tenant_id}/errors/job-{job.id}.xlsx"
    if default_storage.exists(path):
        default_storage.delete(path)
    return default_storage.save(path, ContentFile(payload.getvalue()))


def _parse_excel(job, claim_token: str):
    """Read V1 XLSX, validate every row and atomically refresh preview staging."""

    from hr10_development.legacy.import_job import HrDevelopmentImportJob
    from hr10_development.legacy.staging import HrDevelopmentStagingRow

    if job.template_version != SUPPORTED_TEMPLATE_VERSION:
        raise ValueError(f"UNSUPPORTED_TEMPLATE_VERSION: {job.template_version}")
    if not job.source_file:
        raise ValueError("SOURCE_FILE_MISSING")
    if _sha256(job.source_file) != job.file_hash:
        raise ValueError("SOURCE_FILE_HASH_MISMATCH")
    _heartbeat(job.id, claim_token)

    job.source_file.open("rb")
    try:
        workbook = load_workbook(job.source_file, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        raw_headers = next(rows, None)
        if not raw_headers:
            raise ValueError("EMPTY_WORKBOOK")
        headers = [_normalise_header(value) for value in raw_headers]
        if any(not header for header in headers) or len(headers) != len(set(headers)):
            raise ValueError("INVALID_OR_DUPLICATE_HEADERS")

        schema = TEMPLATE_SCHEMAS[job.job_type]
        missing_headers = sorted(set(schema["required"]) - set(headers))
        if missing_headers:
            raise ValueError(f"MISSING_HEADERS: {','.join(missing_headers)}")

        valid_rows = []
        error_rows = []
        seen_keys = set()
        total = 0
        for row_number, values in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            total += 1
            if total > MAX_IMPORT_ROWS:
                raise ValueError(f"ROW_LIMIT_EXCEEDED: {MAX_IMPORT_ROWS}")
            if total % HEARTBEAT_EVERY_ROWS == 0:
                _heartbeat(job.id, claim_token)

            raw = {headers[index]: _json_value(value) for index, value in enumerate(values) if index < len(headers)}
            parsed, errors = _coerce_row(raw, schema)
            business_key = str(parsed.get(schema["business_key"], "")).strip()
            if business_key and business_key in seen_keys:
                errors.append(f"{schema['business_key']}: DUPLICATE_IN_WORKBOOK")
            seen_keys.add(business_key)
            if errors:
                error_rows.append({"rowNumber": row_number, "raw": raw, "errors": errors})
            else:
                valid_rows.append((row_number, raw, parsed))
    finally:
        job.source_file.close()

    authority_errors = _batch_authority_validation_errors(job.job_type, job.tenant_id, valid_rows)
    authority_valid_rows = []
    for row_number, raw, parsed in valid_rows:
        row_errors = authority_errors.get(row_number) or []
        if row_errors:
            error_rows.append({"rowNumber": row_number, "raw": raw, "errors": row_errors})
        else:
            authority_valid_rows.append((row_number, raw, parsed))
    valid_rows = authority_valid_rows
    error_rows.sort(key=lambda item: item["rowNumber"])

    _heartbeat(job.id, claim_token)
    error_path = _save_error_workbook(job, headers, error_rows)
    now = timezone.now()
    with transaction.atomic():
        locked = HrDevelopmentImportJob.objects.select_for_update().get(id=job.id)
        _require_claim_locked(locked, claim_token, statuses=_PARSE_ACTIVE)
        locked.heartbeat_at = now
        locked.lease_expires_at = _lease_deadline(now)

        HrDevelopmentStagingRow.objects.filter(
            tenant_id=job.tenant_id,
            import_job_id=job.id,
            source_system="EXCEL",
        ).delete()
        HrDevelopmentStagingRow.objects.bulk_create(
            [
                HrDevelopmentStagingRow(
                    tenant_id=job.tenant_id,
                    source_system="EXCEL",
                    source_table=job.job_type,
                    source_field="ROW",
                    source_object_id=f"{job.template_version}:{row_number}",
                    raw_text=json.dumps(raw, ensure_ascii=False, sort_keys=True),
                    parsed_data=parsed,
                    migration_trust_level="DECLARED",
                    target_model=schema["target_model"],
                    import_job_id=job.id,
                    verification_status="PENDING",
                    execution_status="PENDING",
                )
                for row_number, raw, parsed in valid_rows
            ],
            batch_size=500,
        )

        locked.total_rows = total
        locked.processed_rows = len(valid_rows)
        locked.error_rows = len(error_rows)
        locked.warning_rows = 0
        locked.checkpoint_row = total + 1 if total else 1
        locked.error_workbook_path = error_path
        locked.result_summary_json = {
            **dict(locked.result_summary_json or {}),
            "templateVersion": job.template_version,
            "sourceHash": job.file_hash,
            "totalRows": total,
            "validRows": len(valid_rows),
            "errorRows": len(error_rows),
            "previewReady": not error_rows,
            "replaySafe": True,
            "claimLeaseSafe": True,
        }
        locked.save(
            update_fields=[
                "total_rows",
                "processed_rows",
                "error_rows",
                "warning_rows",
                "checkpoint_row",
                "error_workbook_path",
                "result_summary_json",
                "heartbeat_at",
                "lease_expires_at",
                "updated_at",
            ]
        )


def _claim_execution_job(job_id: int, tenant_id: int):
    from hr10_development.legacy.import_job import HrDevelopmentImportJob

    now = timezone.now()
    token = uuid.uuid4().hex
    with transaction.atomic():
        job = HrDevelopmentImportJob.objects.select_for_update().filter(id=job_id, tenant_id=tenant_id).first()
        if not job:
            raise ImportExecutionError("IMPORT_NOT_FOUND", "导入任务不存在")
        if job.status == "SUCCESS":
            return job, None
        if job.status in {"FAILED", "CANCELLED"}:
            raise ImportExecutionError("IMPORT_NOT_READY", f"任务当前状态为 {job.status}")
        if job.error_rows:
            raise ImportExecutionError("IMPORT_HAS_ERRORS", "请修复错误行后重新上传")
        if job.status in _EXECUTE_ACTIVE and _lease_is_active(job, now):
            raise ImportJobBusy(job)
        if job.status not in {"PREVIEW", *_EXECUTE_ACTIVE}:
            raise ImportExecutionError("IMPORT_NOT_READY", "必须先完成校验并进入 PREVIEW")

        takeover = job.status in _EXECUTE_ACTIVE
        summary = dict(job.result_summary_json or {})
        if takeover:
            summary["executionLeaseTakeovers"] = int(summary.get("executionLeaseTakeovers", 0)) + 1

        job.status = "CONFIRMING"
        job.claim_token = token
        job.heartbeat_at = now
        job.lease_expires_at = _lease_deadline(now)
        job.completed_at = None
        job.result_summary_json = summary
        job.save(
            update_fields=[
                "status",
                "claim_token",
                "heartbeat_at",
                "lease_expires_at",
                "completed_at",
                "result_summary_json",
                "updated_at",
            ]
        )
        return job, token


def _mark_executing(job_id: int, tenant_id: int, claim_token: str):
    from hr10_development.legacy.import_job import HrDevelopmentImportJob

    now = timezone.now()
    with transaction.atomic():
        job = HrDevelopmentImportJob.objects.select_for_update().get(id=job_id, tenant_id=tenant_id)
        _require_claim_locked(job, claim_token, statuses={"CONFIRMING", "EXECUTING"})
        job.status = "EXECUTING"
        job.heartbeat_at = now
        job.lease_expires_at = _lease_deadline(now)
        job.save(update_fields=["status", "heartbeat_at", "lease_expires_at", "updated_at"])
        return job


def _build_target(job_type: str, tenant_id: int, parsed: dict, actor_id: int | None):
    from hr10_development.models import HrDevelopmentPlan, HrEnterprisePracticeProject, HrLearningProgram

    common = {"tenant_id": tenant_id}
    if actor_id is not None:
        common.update({"created_by_id": actor_id, "updated_by_id": actor_id})

    if job_type == "EXCEL_PLAN":
        allowed = {"plan_no", "plan_type", "owner_org_id", "start_date", "end_date", "cycle_type"}
        values = {key: parsed[key] for key in allowed if key in parsed}
        if parsed.get("plan_type") == "INDIVIDUAL":
            staff_uuid = parsed.get("staff_master_uuid")
            if not staff_uuid:
                raise ImportExecutionError(
                    "STAFF_IDENTITY_NOT_NORMALIZED",
                    "个人发展计划缺少经 HR03 核验的 canonical UUID",
                )
            values["staff_master_uuid"] = staff_uuid
            values["staff_master_id"] = parsed.get("staff_master_legacy_id")
        return HrDevelopmentPlan(**common, **values)
    if job_type == "EXCEL_PROGRAM":
        allowed = {"program_code", "title", "activity_type", "owner_org_id", "provider_org_id"}
        values = {key: parsed[key] for key in allowed if key in parsed}
        return HrLearningProgram(**common, source="IMPORT", **values)
    if job_type == "EXCEL_PRACTICE":
        allowed = {
            "project_no",
            "title",
            "provider_org_id",
            "owner_org_id",
            "planned_start_date",
            "planned_end_date",
            "capacity",
        }
        values = {key: parsed[key] for key in allowed if key in parsed}
        return HrEnterprisePracticeProject(**common, **values)
    raise ImportExecutionError("UNSUPPORTED_JOB_TYPE", f"不支持的导入类型: {job_type}")


def _row_number(source_object_id: str) -> int | None:
    try:
        return int(str(source_object_id).rsplit(":", 1)[-1])
    except (TypeError, ValueError):
        return None


def _execute_rows_atomic(job, claim_token: str, *, actor_id: int | None, request_id: str, reason: str):
    from hr10_development.legacy.import_job import HrDevelopmentImportJob
    from hr10_development.legacy.staging import HrDevelopmentStagingRow
    from hr10_development.models import HrDevelopmentAuditEvent

    now = timezone.now()
    schema = TEMPLATE_SCHEMAS.get(job.job_type)
    if not schema:
        raise ImportExecutionError("UNSUPPORTED_JOB_TYPE", f"不支持的导入类型: {job.job_type}")

    with transaction.atomic():
        locked = HrDevelopmentImportJob.objects.select_for_update().get(id=job.id, tenant_id=job.tenant_id)
        _require_claim_locked(locked, claim_token, statuses={"EXECUTING"})
        locked.heartbeat_at = now
        locked.lease_expires_at = _lease_deadline(now)

        rows = list(
            HrDevelopmentStagingRow.objects.select_for_update()
            .filter(tenant_id=job.tenant_id, import_job_id=job.id, source_system="EXCEL")
            .order_by("id")
        )
        if len(rows) != locked.processed_rows:
            raise ImportExecutionError("STAGING_ROW_COUNT_MISMATCH", "预览暂存行数量与任务统计不一致")

        execution_rows = []
        for row in rows:
            row_no = _row_number(row.source_object_id)
            if row.source_table != job.job_type or row.target_model != schema["target_model"]:
                raise ImportExecutionError(
                    "STAGING_TARGET_MISMATCH",
                    "暂存行目标模型与导入任务不一致",
                    row_id=row.id,
                    row_number=row_no,
                )
            if row.execution_status != "PENDING" or row.target_id is not None:
                raise ImportExecutionError(
                    "STAGING_ROW_ALREADY_EXECUTED",
                    "暂存行执行状态异常，拒绝重复写入",
                    row_id=row.id,
                    row_number=row_no,
                )
            execution_rows.append((row_no, {}, dict(row.parsed_data or {})))

        # Re-check all authority keys/references immediately before the write, but do
        # it in batches so a 10k-row confirmation cannot degenerate into N+1 queries.
        authority_errors = _batch_authority_validation_errors(job.job_type, job.tenant_id, execution_rows)
        for row, (row_no, _, _) in zip(rows, execution_rows):
            current_errors = authority_errors.get(row_no) or []
            if current_errors:
                raise ImportExecutionError(
                    "AUTHORITY_CONFLICT_AT_EXECUTION",
                    "; ".join(current_errors),
                    row_id=row.id,
                    row_number=row_no,
                )

        audit_events = []
        imported = 0
        for row, (row_no, _, parsed_data) in zip(rows, execution_rows):
            target = _build_target(job.job_type, job.tenant_id, parsed_data, actor_id)
            try:
                target.full_clean()
                target.save(force_insert=True)
            except (ValidationError, IntegrityError) as exc:
                raise ImportExecutionError(
                    "TARGET_WRITE_REJECTED",
                    str(exc),
                    row_id=row.id,
                    row_number=row_no,
                ) from exc

            row.target_id = target.pk
            row.execution_status = "SUCCESS"
            row.executed_at = now
            row.error_message = ""
            row.updated_by_id = actor_id
            row.save(
                update_fields=[
                    "target_id",
                    "execution_status",
                    "executed_at",
                    "error_message",
                    "updated_by",
                    "updated_at",
                ]
            )
            imported += 1

            audit_events.append(
                HrDevelopmentAuditEvent(
                    tenant_id=job.tenant_id,
                    actor_id_id=actor_id,
                    object_type=schema["target_model"],
                    object_id=str(target.pk),
                    action="ImportRowExecuted",
                    after_json={
                        "importJobId": str(job.id),
                        "sourceRow": row_no,
                        "businessKey": row.parsed_data.get(schema["business_key"]),
                        "targetModel": schema["target_model"],
                    },
                    reason=reason[:1000],
                    request_id=request_id[:64],
                )
            )

        completed_at = timezone.now()
        summary = dict(locked.result_summary_json or {})
        summary.update(
            {
                "confirmed": True,
                "confirmedAt": completed_at.isoformat(),
                "executedRows": imported,
                "executionFailedRows": 0,
                "allOrNothing": True,
                "auditEvents": imported + 1,
            }
        )
        locked.status = "SUCCESS"
        locked.completed_at = completed_at
        locked.claim_token = ""
        locked.lease_expires_at = None
        locked.heartbeat_at = completed_at
        locked.result_summary_json = summary
        if actor_id is not None:
            locked.updated_by_id = actor_id
        locked.save(
            update_fields=[
                "status",
                "completed_at",
                "claim_token",
                "lease_expires_at",
                "heartbeat_at",
                "result_summary_json",
                "updated_by",
                "updated_at",
            ]
        )

        audit_events.append(
            HrDevelopmentAuditEvent(
                tenant_id=job.tenant_id,
                actor_id_id=actor_id,
                object_type="HrDevelopmentImportJob",
                object_id=str(job.id),
                action="ImportConfirmedAndExecuted",
                after_json={
                    "jobType": job.job_type,
                    "executedRows": imported,
                    "status": "SUCCESS",
                },
                reason=reason[:1000],
                request_id=request_id[:64],
            )
        )
        HrDevelopmentAuditEvent.objects.bulk_create(audit_events, batch_size=500)
        return locked


def _mark_execution_failed(job_id: int, tenant_id: int, claim_token: str, exc: Exception, *, actor_id: int | None, request_id: str, reason: str):
    from hr10_development.legacy.import_job import HrDevelopmentImportJob
    from hr10_development.legacy.staging import HrDevelopmentStagingRow
    from hr10_development.models import HrDevelopmentAuditEvent

    now = timezone.now()
    with transaction.atomic():
        job = HrDevelopmentImportJob.objects.select_for_update().filter(id=job_id, tenant_id=tenant_id).first()
        if not job or job.claim_token != claim_token:
            return job

        if isinstance(exc, ImportExecutionError) and exc.row_id:
            HrDevelopmentStagingRow.objects.filter(
                id=exc.row_id,
                tenant_id=tenant_id,
                import_job_id=job_id,
            ).update(
                execution_status="FAILED",
                error_message=str(exc)[:2000],
                executed_at=now,
                updated_by_id=actor_id,
            )

        summary = dict(job.result_summary_json or {})
        summary.update(
            {
                "confirmed": True,
                "errorCode": getattr(exc, "code", "IMPORT_EXECUTION_FAILED"),
                "error": str(exc)[:2000],
                "failedPhase": "EXECUTE",
                "failedRow": getattr(exc, "row_number", None),
                "allOrNothing": True,
                "rolledBack": True,
            }
        )
        job.status = "FAILED"
        job.completed_at = now
        job.claim_token = ""
        job.lease_expires_at = None
        job.heartbeat_at = now
        job.result_summary_json = summary
        if actor_id is not None:
            job.updated_by_id = actor_id
        job.save(
            update_fields=[
                "status",
                "completed_at",
                "claim_token",
                "lease_expires_at",
                "heartbeat_at",
                "result_summary_json",
                "updated_by",
                "updated_at",
            ]
        )
        HrDevelopmentAuditEvent.objects.create(
            tenant_id=tenant_id,
            actor_id_id=actor_id,
            object_type="HrDevelopmentImportJob",
            object_id=str(job_id),
            action="ImportExecutionFailed",
            after_json={
                "status": "FAILED",
                "errorCode": summary["errorCode"],
                "failedRow": summary.get("failedRow"),
                "rolledBack": True,
            },
            reason=reason[:1000],
            request_id=request_id[:64],
        )
        return job


def execute_import_job(
    job_id: int,
    *,
    tenant_id: int,
    actor_id: int | None = None,
    request_id: str = "",
    reason: str = "Excel import confirmed",
):
    """Execute a PREVIEW job into authority tables with one transactional commit.

    SUCCESS is written only in the same transaction as all target rows and all
    success audit records. A process crash leaves the EXECUTING lease to expire;
    a later confirmation can safely reclaim it because uncommitted writes roll
    back with the database connection.
    """

    claim_token = None
    try:
        job, claim_token = _claim_execution_job(job_id, tenant_id)
        if not claim_token:
            return job
        job = _mark_executing(job_id, tenant_id, claim_token)
        return _execute_rows_atomic(
            job,
            claim_token,
            actor_id=actor_id,
            request_id=request_id,
            reason=reason,
        )
    except ImportJobBusy:
        raise
    except ImportExecutionError as exc:
        if claim_token:
            _mark_execution_failed(
                job_id,
                tenant_id,
                claim_token,
                exc,
                actor_id=actor_id,
                request_id=request_id,
                reason=reason,
            )
        raise
    except Exception as exc:
        logger.exception("Import job %s execution failed", job_id)
        wrapped = ImportExecutionError("IMPORT_EXECUTION_FAILED", str(exc)[:2000])
        if claim_token:
            _mark_execution_failed(
                job_id,
                tenant_id,
                claim_token,
                wrapped,
                actor_id=actor_id,
                request_id=request_id,
                reason=reason,
            )
        raise wrapped from exc
