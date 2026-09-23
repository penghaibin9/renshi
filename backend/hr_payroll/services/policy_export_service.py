"""Human-readable, scoped export of an immutable payroll trial, not a bank file."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .policy_import_service import safe_text
from .policy_math import decimal

def trial_workbook(trial):
    from .policy_payroll_service import verify_trial
    verify_trial(trial)
    wb=Workbook();summary=wb.active;summary.title='试算回执'
    data=trial.input_payload_json;out=trial.output_json
    summary.append(['说明','试算与复核依据，不是已发工资或银行付款文件'])
    for name,value in [('试算编号',str(trial.id)),('修订',trial.revision_no),('学校编号',str(trial.tenant_id)),('人员编号',str(trial.staff_id)),('用途',trial.purpose),('所属期间',data.get('periodCode','历史补差')),('预计支付日',data.get('paymentDate')),('内容校验',trial.content_hash)]:summary.append([name,safe_text(value)])
    for name,key in [('应发','gross'),('个人扣款','deduction'),('实发','net'),('单位成本','employerCost'),('总成本','totalCost')]:summary.append([name,decimal(out[key])])
    lines=wb.create_sheet('逐项金额');lines.append(['工资项目','名称','性质','金额','规则编号'])
    segments=wb.create_sheet('分段依据');segments.append(['工资项目','生效日','失效日不含当日','计付天数','全月天数','计付方法','规则版本','个人核定版本','制度版本'])
    for item in out['lines']:
        lines.append([safe_text(item['itemCode']),safe_text(item['name']),{'EARNING':'应发','DEDUCTION':'个人扣款','EMPLOYER':'单位成本'}.get(item['type'],item['type']),decimal(item['amount']),safe_text(item.get('ruleId',''))])
        for row in item.get('segments',[]):segments.append([safe_text(item['itemCode']),row['from'],row['toExclusive'],row['days'],row['monthDays'],row['allocation'],row['ruleId'],row['basisId'],row['policyId']])
    cost=wb.create_sheet('成本分摊');cost.append(['经费或部门代码','分摊金额'])
    for key,value in out.get('costAllocation',{}).items():cost.append([safe_text(key),decimal(value)])
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245B92')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=30
        for row in ws:
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.data_type=='n':cell.number_format='#,##0.00;[Red](#,##0.00);"-"'
    summary.column_dimensions['B'].width=75
    stream=BytesIO();wb.save(stream);return stream.getvalue()
