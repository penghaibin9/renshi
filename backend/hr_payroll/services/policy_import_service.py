"""Bounded xlsx -> inspectable stage -> atomic draft creation. Never auto-publish."""
from io import BytesIO
from datetime import date, datetime, timedelta
from decimal import Decimal
import hashlib
import json
import zipfile
from django.db import transaction, IntegrityError
from django.core.exceptions import ValidationError
from django.utils import timezone
from hr_payroll.policy_models import PayrollImportStage
from .policy_configuration_service import PolicyConfigurationService, actor_required
from .policy_math import PolicyPayrollError, digest

HEADERS = {
    "STANDARD": ["payGroupCode", "tableCode", "levelCode", "amount", "currencyCode", "effectiveFrom", "effectiveTo", "evidenceRef", "supersedesId"],
    "BASIS": ["payrollProfileId", "selectors", "variables", "taxDeductions", "costShares", "effectiveFrom", "effectiveTo", "evidenceRef", "supersedesId"],
    "WORKLOAD": ["periodId", "staffId", "workloadKey", "variableKey", "units", "share", "coefficient", "effectiveFrom", "effectiveTo", "evidenceRef", "supersedesId"],
}
JSON_FIELDS = {"selectors", "variables", "taxDeductions", "costShares"}
NUM_FIELDS = {"amount", "units", "share", "coefficient"}


def safe_text(value):
    text = str("" if value is None else value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def workbook_bytes(kind, rows=(), errors=()):
    """Application-side download; strings are escaped against spreadsheet formula injection."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "导入数据"
    ws.append(HEADERS[kind])
    for row in rows:
        ws.append([safe_text(json.dumps(row.get(k), ensure_ascii=False)) if isinstance(row.get(k), dict) else safe_text(row.get(k, "")) for k in HEADERS[kind]])
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="245B92")
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 24
    help_ws=wb.create_sheet("填写说明")
    help_ws.append(["字段名必须保持不变；不填示例金额，不含任何全国通用工资标准。"])
    help_ws.append(["日期为YYYY-MM-DD，生效区间左闭右开；金额为十进制；不得含公式。"])
    help_ws.append(["BASIS的selectors/variables/taxDeductions/costShares填写JSON；推荐在页面表单核定后导入。"])
    help_ws.append(["上传只预览；确认仅生成草稿；另一位授权人员复核发布后才用于试算。"])
    help_ws.column_dimensions["A"].width=108
    if errors:
        er=wb.create_sheet("错误行");er.append(["Excel行号","错误代码","处理提示"])
        for x in errors: er.append([x["row"],safe_text(x["code"]),safe_text(x["message"])])
        for key,width in [("A",16),("B",46),("C",85)]:er.column_dimensions[key].width=width
        er.freeze_panes="A2"
    buf=BytesIO();wb.save(buf);return buf.getvalue()


class PayrollPolicyImportService:
    def __init__(self,tenant_id,actor_user_id):
        self.tenant=int(tenant_id);self.actor=actor_required(actor_user_id)
        if self.tenant<=0:raise PolicyPayrollError("TENANT_CONTEXT_REQUIRED","学校上下文无效")

    def preview(self,kind,blob):
        from openpyxl import load_workbook
        kind=kind.upper()
        if kind not in HEADERS:raise PolicyPayrollError("PAYROLL_IMPORT_KIND_INVALID","仅支持标准、核定、工作量导入")
        if not blob or len(blob)>8*1024*1024:raise PolicyPayrollError("PAYROLL_IMPORT_SIZE_LIMIT","仅接受不超过8 MiB的xlsx文件")
        try:
            with zipfile.ZipFile(BytesIO(blob)) as z:
                items=z.infolist()
                if (len(items)>512 or sum(x.file_size for x in items)>40*1024*1024
                        or any(x.flag_bits & 1 or x.filename.endswith("vbaProject.bin") or "externalLinks/" in x.filename for x in items)):
                    raise PolicyPayrollError("PAYROLL_IMPORT_ARCHIVE_LIMIT","压缩展开体积、文件数量或宏不符合导入要求")
                if any(x.file_size>1024*1024 and x.file_size>max(x.compress_size,1)*250 for x in items):
                    raise PolicyPayrollError("PAYROLL_IMPORT_ARCHIVE_LIMIT","工作簿压缩比异常")
            wb=load_workbook(BytesIO(blob),read_only=True,data_only=False,keep_links=False)
        except (zipfile.BadZipFile,KeyError,ValueError,OSError) as exc:
            raise PolicyPayrollError("PAYROLL_IMPORT_FILE_INVALID","文件不是可读取的xlsx工作簿") from exc
        rows=[];errors=[]
        try:
            if "导入数据" not in wb.sheetnames:raise PolicyPayrollError("PAYROLL_IMPORT_SHEET_REQUIRED","缺少导入数据工作表")
            ws=wb["导入数据"]
            if ws.max_row>1001 or ws.max_column>len(HEADERS[kind]):raise PolicyPayrollError("PAYROLL_IMPORT_ROW_LIMIT","一次最多1000行，请按模板列填写")
            iterator=ws.iter_rows()
            first=next(iterator,[])
            if [c.value for c in first]!=HEADERS[kind]:raise PolicyPayrollError("PAYROLL_IMPORT_HEADER_INVALID","字段名或顺序不符，请使用本页模板")
            seen=set()
            for i,cells in enumerate(iterator,2):
                if i > 1001: raise PolicyPayrollError("PAYROLL_IMPORT_ROW_LIMIT", "一次最多1000行")
                if all(c.value is None for c in cells):continue
                row={};bad=False
                try:
                    for key,cell in zip(HEADERS[kind],cells):
                        value=cell.value
                        if cell.data_type=="f" or cell.data_type=="e":raise PolicyPayrollError("PAYROLL_IMPORT_FORMULA_FORBIDDEN","请粘贴确定的数值，不接受公式或Excel错误值")
                        if isinstance(value,str) and len(value)>16000:raise PolicyPayrollError("PAYROLL_IMPORT_CELL_LIMIT","单元格内容过长")
                        if value is None or value=="":continue
                        if key in JSON_FIELDS:
                            value=json.loads(str(value));
                            if not isinstance(value,dict):raise ValueError("object required")
                        elif key in NUM_FIELDS:value=str(value)
                        elif isinstance(value,(date,datetime)):value=value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
                        else:value=str(value).strip()
                        row[key]=value
                    identity=tuple(str(row.get(k,"")) for k in ({"STANDARD":["payGroupCode","tableCode","levelCode","effectiveFrom"],"BASIS":["payrollProfileId","effectiveFrom"],"WORKLOAD":["periodId","workloadKey","staffId","effectiveFrom"]}[kind]))
                    if identity in seen:raise PolicyPayrollError("PAYROLL_IMPORT_DUPLICATE_ROW","同一文件存在重复业务行")
                    seen.add(identity)
                    # The real domain validation is used; savepoint is always rolled back.
                    with transaction.atomic():
                        PolicyConfigurationService(self.tenant,self.actor).create(kind,row)
                        transaction.set_rollback(True)
                except (PolicyPayrollError,ValidationError,ValueError,TypeError,IntegrityError) as exc:
                    errors.append({"row":i,"code":getattr(exc,"code","PAYROLL_IMPORT_ROW_INVALID"),"message":str(exc)[:500] if isinstance(exc,PolicyPayrollError) else "字段、日期、金额或本校关联记录无效；请核对本行"})
                rows.append({"excelRow":i,"data":row})
        finally:wb.close()
        if not rows:raise PolicyPayrollError("PAYROLL_IMPORT_EMPTY","模板没有业务数据")
        file_hash=hashlib.sha256(blob).hexdigest()
        payload={"tenantId":self.tenant,"kind":kind,"fileHash":file_hash,"rows":rows,"errors":errors}
        return PayrollImportStage.objects.create(tenant_id=self.tenant,created_by=self.actor,updated_by=self.actor,
               kind=kind,file_hash=file_hash,stage_hash=digest(payload),rows_json=rows,errors_json=errors,
               expires_at=timezone.now()+timedelta(hours=24))

    @transaction.atomic
    def confirm(self,stage_id,expected_hash):
        stage=PayrollImportStage.objects.select_for_update().filter(tenant_id=self.tenant,id=stage_id,created_by=self.actor).first()
        if not stage:raise PolicyPayrollError("PAYROLL_IMPORT_NOT_FOUND","只能确认本人在本校预览的导入")
        payload={"tenantId":self.tenant,"kind":stage.kind,"fileHash":stage.file_hash,"rows":stage.rows_json,"errors":stage.errors_json}
        if stage.stage_hash!=expected_hash or digest(payload)!=stage.stage_hash:raise PolicyPayrollError("PAYROLL_IMPORT_CHANGED","预览内容已变化，请重新预览")
        if stage.status=="APPLIED":return stage
        if stage.expires_at<=timezone.now():raise PolicyPayrollError("PAYROLL_IMPORT_EXPIRED","预览超过24小时，请重新核对")
        if stage.errors_json:raise PolicyPayrollError("PAYROLL_IMPORT_ERRORS_REMAIN","有错误行，不允许部分静默导入")
        from hr_payroll.policy_models import PayrollPolicyScope
        scope, _ = PayrollPolicyScope.objects.get_or_create(tenant_id=self.tenant, scope_key="IMPORT:"+stage.file_hash)
        PayrollPolicyScope.objects.select_for_update().get(pk=scope.pk)
        if PayrollImportStage.objects.filter(tenant_id=self.tenant,kind=stage.kind,file_hash=stage.file_hash,status="APPLIED").exclude(id=stage.id).exists():
            raise PolicyPayrollError("PAYROLL_IMPORT_ALREADY_APPLIED", "同一文件已确认过，请查阅原草稿，不重复创建")
        ids=[]
        for row in stage.rows_json:
            obj=PolicyConfigurationService(self.tenant,self.actor).create(stage.kind,row["data"]);ids.append(str(obj.id))
        stage.applied_ids_json=ids;stage.status="APPLIED";stage.updated_by=self.actor
        stage.save(update_fields=["applied_ids_json","status","updated_by","updated_at"])
        from horilla.hr_event_service import emit_registered_event
        emit_registered_event(tenant_id=self.tenant,event_name="hr.payroll.policy.imported",payload={"stageId":str(stage.id),"stageHash":stage.stage_hash,"count":len(ids),"actorId":self.actor},correlation_id="")
        return stage
