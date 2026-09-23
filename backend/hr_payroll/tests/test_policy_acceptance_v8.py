"""Acceptance edges on actual ORM; company membership is an explicit boundary fixture.

These are not claims about a real bank, Windows desktop or MySQL locking.
"""
import copy
import json
from datetime import date
from decimal import Decimal
from unittest.mock import patch
from django.test import TestCase, RequestFactory, override_settings
from hr_payroll.tests import test_policy_payroll_depth as seed
from hr_payroll.policy_models import PayrollTaxReservation, PayrollTaxAccount, PayrollTaxReceiptSupplement, PayrollTrial
from hr_payroll.services.policy_tax_service import PolicyTaxService
from hr_payroll.services.policy_math import PolicyPayrollError
from hr_payroll.services.payment_service import PayrollPaymentService, PayrollPaymentError
from hr_payroll.services.statutory_contribution_service import StatutoryContributionRuleService
from hr_payroll import policy_api
from hr_payroll.authority_registry import PERM_REVIEW, PERM_RULE_MANAGE, PERM_CALCULATE

class Actor:
    is_authenticated=True
    is_superuser=False
    def __init__(self,ident,perms):self.id=ident;self.permissions=set(perms)|{'hr.payroll.view'}
    def has_perm(self,p):return p in self.permissions

@override_settings(HR15_PAYROLL_INPUT_PROVIDERS=seed.PROVIDERS,HR15_PAYMENT_PROVIDERS=seed.BANK)
class PolicyAcceptanceV8Tests(TestCase):
    tenant=701
    setUp=seed.PayrollPolicyOrmTests.setUp
    person=seed.PayrollPolicyOrmTests.person
    publish=seed.PayrollPolicyOrmTests.publish
    approved_trial=seed.PayrollPolicyOrmTests.approved_trial
    calculated=seed.PayrollPolicyOrmTests.calculated
    paid=seed.PayrollPolicyOrmTests.paid

    def test_missing_date_cannot_be_fully_reconciled(self):
        result,instruction,receipt=self.paid(paid_date=None)
        self.assertEqual(PayrollTaxReservation.objects.get().status,'DATE_REVIEW')
        self.assertEqual(PayrollTaxAccount.objects.get().version_no,0)
        with self.assertRaises(PayrollPaymentError) as c:
            PayrollPaymentService(self.tenant,103).reconcile(instruction_id=instruction.id,reconciliation_no='TEST-R')
        self.assertEqual(c.exception.code,'PAYROLL_TAX_RECONCILIATION_PENDING')
        instruction.refresh_from_db();self.assertEqual(instruction.status,'ACCEPTED')

    def test_verified_date_supplement_is_append_only_and_idempotent(self):
        result,instruction,receipt=self.paid(paid_date=None)
        instruction.refresh_from_db();original=copy.deepcopy(instruction.provider_receipt_json)
        service=PolicyTaxService(self.tenant,103)
        proof={**receipt,'paidDate':'2026-09-30'}
        service.supplement_verified_date(instruction_id=instruction.id,provider_payload=proof)
        service.supplement_verified_date(instruction_id=instruction.id,provider_payload=proof)
        instruction.refresh_from_db()
        self.assertEqual(instruction.provider_receipt_json,original)
        self.assertEqual(PayrollTaxReceiptSupplement.objects.count(),1)
        self.assertEqual(PayrollTaxAccount.objects.get().version_no,1)
        recon=PayrollPaymentService(self.tenant,103).reconcile(instruction_id=instruction.id,reconciliation_no='TEST-R')
        self.assertEqual(recon.status,'MATCHED')

    def test_cross_month_receipt_needs_correction_not_fake_reconciliation(self):
        _,instruction,receipt=self.paid(paid_date=None)
        with self.assertRaises(PolicyPayrollError) as c:
            PolicyTaxService(self.tenant,103).supplement_verified_date(instruction_id=instruction.id,provider_payload={**receipt,'paidDate':'2026-10-01'})
        self.assertEqual(c.exception.code,'PAYROLL_TAX_CROSS_MONTH_CORRECTION_REQUIRED')
        self.assertEqual(PayrollTaxReceiptSupplement.objects.count(),0)
        self.assertEqual(PayrollTaxAccount.objects.get().version_no,0)

    def test_changed_receipt_identity_is_rejected(self):
        _,instruction,receipt=self.paid(paid_date=None)
        with self.assertRaises((PolicyPayrollError,PayrollPaymentError)):
            PolicyTaxService(self.tenant,103).supplement_verified_date(instruction_id=instruction.id,provider_payload={**receipt,'receiptNo':'OTHER','paidDate':'2026-09-30'})
        self.assertEqual(PayrollTaxReceiptSupplement.objects.count(),0)

    def test_statutory_floor_and_employer_cost_do_not_double_deduct(self):
        rules=StatutoryContributionRuleService(self.tenant,101)
        rule=rules.create_draft(rule_code='TEST-PENSION',version_no=1,contribution_group='SOCIAL_INSURANCE',contribution_code='TEST-PENSION',name='仅测试养老',jurisdiction_code='TEST-AREA',base_variable_key='pensionBase',base_floor='4000',base_ceiling='9000',employee_rate='0.08',employer_rate='0.16',employee_item_code='PENSION_PERSONAL',employer_item_code='PENSION_EMPLOYER',effective_from=date(2026,9,1),policy_evidence={'documentNo':'FICTITIOUS-TEST-NOT-POLICY'})
        StatutoryContributionRuleService(self.tenant,102).publish(rule.id)
        cfg=copy.deepcopy(self.policy.configuration_json);cfg.update(statutoryMode='REQUIRED',statutoryCodes=['TEST-PENSION'],statutoryReason='')
        self.publish('POLICY',{'payGroupCode':'FACULTY','name':'测试参保组','configuration':cfg,'supersedesId':str(self.policy.id)})
        self.publish('BASIS',{'payrollProfileId':str(self.profile.id),'selectors':{'positionGrade':'G1','employmentRelationshipId':str(self.relation.id),'statutoryJurisdiction':'TEST-AREA'},'variables':{'pensionBase':'1000'},'taxDeductions':seed.tax_data(),'costShares':{},'supersedesId':str(self.basis.id)})
        trial=self.approved_trial()
        self.assertEqual(Decimal(trial.output_json['deduction']),Decimal('320'))
        self.assertEqual(Decimal(trial.output_json['net']),Decimal('5680'))
        self.assertEqual(Decimal(trial.output_json['employerCost']),Decimal('640'))
        self.assertEqual(Decimal(trial.output_json['totalCost']),Decimal('6640'))

    def test_month_end_exclusive_does_not_drop_last_day(self):
        self.profile.effective_to=date(2026,10,1);self.profile.save()
        self.assertEqual(Decimal(self.approved_trial().output_json['gross']),Decimal('6000'))

    def test_partial_new_starter_is_not_paid_for_prior_days(self):
        self.relation.effective_from=date(2026,9,16);self.relation.save()
        self.assignment.effective_from=date(2026,9,16);self.assignment.save()
        self.profile.effective_from=date(2026,9,16);self.profile.save()
        self.assertEqual(Decimal(self.approved_trial().output_json['gross']),Decimal('3000'))

    def test_stopped_profile_retains_historical_pay(self):
        self.profile.effective_to=date(2026,9,16);self.profile.status='ENDED';self.profile.save()
        self.relation.effective_to=date(2026,9,16);self.relation.save()
        self.assignment.effective_to=date(2026,9,16);self.assignment.save()
        self.assertEqual(Decimal(self.approved_trial().output_json['gross']),Decimal('3000'))

    def call(self,view,actor,body=None,tenant=None,member=True,**kwargs):
        factory=RequestFactory();request=factory.get('/test/') if body is None else factory.post('/test/',data=json.dumps(body),content_type='application/json')
        request.user=actor
        with patch('hr_payroll.api.resolve_tenant_from_request',return_value=tenant or self.tenant),patch('hr_payroll.api.get_allowed_company_ids',return_value={self.tenant} if member else set()):
            return view(request,**kwargs)

    def test_http_read_role_cannot_make_trial(self):
        r=self.call(policy_api.trials,Actor(101,[]),{'periodId':str(self.period.id),'staffId':str(self.staff.id),'idempotencyKey':'HTTP'})
        self.assertEqual(r.status_code,403);self.assertEqual(PayrollTrial.objects.count(),0)

    def test_http_maker_self_review_denied_and_reviewer_succeeds(self):
        r=self.call(policy_api.trials,Actor(101,[PERM_CALCULATE]),{'periodId':str(self.period.id),'staffId':str(self.staff.id),'idempotencyKey':'HTTP'})
        self.assertEqual(r.status_code,201);data=json.loads(r.content)['data']
        body={'expectedHash':data['hash'],'note':'真实表单复核测试'}
        denied=self.call(policy_api.review_trial,Actor(101,[PERM_REVIEW]),body,trial_id=data['id'])
        self.assertEqual(denied.status_code,400)
        accepted=self.call(policy_api.review_trial,Actor(102,[PERM_REVIEW]),body,trial_id=data['id'])
        self.assertEqual(accepted.status_code,200)

    def test_http_extra_amount_cannot_override_rules(self):
        r=self.call(policy_api.trials,Actor(101,[PERM_CALCULATE]),{'periodId':str(self.period.id),'staffId':str(self.staff.id),'idempotencyKey':'HTTP','gross':'1'})
        self.assertEqual(r.status_code,400);self.assertEqual(PayrollTrial.objects.count(),0)

    def test_http_cross_school_membership_and_object_lookup(self):
        trial=self.approved_trial()
        denied=self.call(policy_api.trial_detail,Actor(102,[PERM_REVIEW]),tenant=702,trial_id=trial.id)
        self.assertEqual(denied.status_code,403)
        missing=self.call(policy_api.trial_detail,Actor(102,[PERM_REVIEW]),trial_id='00000000-0000-0000-0000-000000000001')
        self.assertEqual(missing.status_code,404)

    def test_http_rule_manager_cannot_read_personal_basis(self):
        r=self.call(policy_api.configurations,Actor(101,[PERM_RULE_MANAGE]),kind='BASIS')
        self.assertEqual(r.status_code,403)

    def test_normal_wages_cannot_be_disguised_as_annual_bonus(self):
        cfg=copy.deepcopy(self.policy.configuration_json);cfg['taxMethod']='ANNUAL_BONUS_SEPARATE'
        with self.assertRaises(PolicyPayrollError):self.publish('POLICY',{'payGroupCode':'OTHER','name':'invalid','configuration':cfg})

    def test_trial_xlsx_export_contains_amounts_and_audit(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from hr_staff.models import HrOutboxEvent
        trial=self.approved_trial()
        response=self.call(policy_api.export_trial,Actor(102,[PERM_REVIEW]),trial_id=trial.id)
        self.assertEqual(response.status_code,200)
        wb=load_workbook(BytesIO(response.content),data_only=False)
        self.assertIn('分段依据',wb.sheetnames)
        self.assertEqual(wb['逐项金额']['D2'].value,6000)
        self.assertFalse(any(c.data_type=='f' for ws in wb for row in ws for c in row))
        self.assertTrue(HrOutboxEvent.objects.filter(tenant_id=self.tenant,event_type='hr.payroll.trial.exported').exists())

    def test_midmonth_statutory_change_cannot_silently_use_first_rate(self):
        for version,start,end,rate in [(1,date(2026,9,1),date(2026,9,16),'0.08'),(2,date(2026,9,16),None,'0.09')]:
            service=StatutoryContributionRuleService(self.tenant,101)
            rule=service.create_draft(rule_code='TEST-SOC',version_no=version,contribution_group='SOCIAL_INSURANCE',contribution_code='TEST-SOC',name='测试非政策',jurisdiction_code='TEST-AREA',base_variable_key='pensionBase',base_floor='4000',base_ceiling='9000',employee_rate=rate,employer_rate='0.16',employee_item_code='SOC_PERSON',employer_item_code='SOC_UNIT',effective_from=start,effective_to=end,policy_evidence={'documentNo':'TEST-ONLY'})
            StatutoryContributionRuleService(self.tenant,102).publish(rule.id)
        cfg=copy.deepcopy(self.policy.configuration_json);cfg.update(statutoryMode='REQUIRED',statutoryCodes=['TEST-SOC'],statutoryReason='')
        self.publish('POLICY',{'payGroupCode':'FACULTY','name':'测试参保','configuration':cfg,'supersedesId':str(self.policy.id)})
        self.publish('BASIS',{'payrollProfileId':str(self.profile.id),'selectors':{'positionGrade':'G1','employmentRelationshipId':str(self.relation.id),'statutoryJurisdiction':'TEST-AREA'},'variables':{'pensionBase':'4000'},'taxDeductions':seed.tax_data(),'costShares':{},'supersedesId':str(self.basis.id)})
        with self.assertRaises(PolicyPayrollError) as c:self.approved_trial()
        self.assertEqual(c.exception.code,'PAYROLL_STATUTORY_MONTH_CHANGE_REQUIRES_REVIEW')
