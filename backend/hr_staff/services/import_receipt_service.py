"""P3 read-only, row-addressable migration proof. No personal source values.

A committed row is not a verified migration until the same school, source row,
four authority layers and row audit agree. This is historical-source evidence,
not a claim that the employee is currently active or an account was created.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timezone
from io import BytesIO
from uuid import UUID

from django.db import transaction
from hr_staff.models import (HrImportIssue, HrImportJob, HrStaffMaster, HrEmploymentRelationship,
    HrStaffAssignment, HrStaffAuditEvent)

MAX_RECEIPT_ROWS = 5000
MAX_RECEIPT_ISSUES = 100_000


class ImportReceiptLimit(ValueError):
    """A bounded read cannot prove the complete historical ledger."""

ISSUE_LABELS = {
    "ROW_TENANT_MISMATCH":"行的学校归属不一致，请管理员核查",
    "LEGACY_OR_INVALID_RECEIPT":"历史回执缺失或编号损坏，不能当作迁入成功",
    "PERSON_OR_STAFF_MISSING":"自然人或本校主档缺失，核对原提交回执",
    "DUPLICATE_STAFF_RECEIPT":"多个导入行指向同一主档，请核对重复来源",
    "EMPLOYMENT_SOURCE_MISMATCH":"用工关系与导入行来源不一致",
    "ASSIGNMENT_SOURCE_MISMATCH":"主任职与用工关系、导入行来源不一致",
    "ROW_AUDIT_MISSING_OR_DUPLICATE":"逐行审计缺失或重复，需管理员核对",
    "SOURCE_ROW_REQUIRES_CORRECTION":"本行未写入，请按原文件行号修正后另批预览",
    "COMMITTED_ROW_INVALID":"已提交行被标记无效，需核对原记录和审计",
    "ROW_STATE_INVALID":"历史行状态不符合导入约定，不能自动判定成功",
    "ASSIGNMENT_REFERENCE_SCOPE_MISMATCH":"任职引用的部门或岗位不属于本校",
    "ROW_NOT_COMMITTED":"本行尚未完成提交，先查原任务，不重复上传",
}


def safe_issue_rows(job):
    from hr_staff.import_tabular import COLUMNS
    # Do not trust historical free-text diagnostics to be free of personal data.
    # Export an actionable category and source row, never the raw source values.
    labels={"VALIDATION_ERROR":"字段预检未通过；在原任务中核对字段要求",
        "COMMIT_VALIDATION_ERROR":"提交时字段或部门依据变化，需重新核对",
        "COMMIT_FAILED":"提交失败，原事务已回滚；核对原任务错误后再办理"}
    records = list(HrImportIssue.objects.filter(tenant_id=job.tenant_id, job_id=job)
        .order_by("row_no", "field_code", "pk")
        .values("row_no", "field_code", "error_code")[:MAX_RECEIPT_ISSUES + 1])
    if len(records) > MAX_RECEIPT_ISSUES:
        raise ImportReceiptLimit("IMPORT_RECEIPT_ISSUE_LIMIT")
    return [[x["row_no"], COLUMNS.get(x["field_code"], "整行"),
        labels.get(x["error_code"], "需在原任务中核对该行反馈"),
        x["error_code"] if x["error_code"] in labels else "UNCLASSIFIED_ERROR"]
        for x in records]



def _uuid(value):
    try:
        return str(UUID(str(value)))
    except (ValueError, TypeError, AttributeError):
        return None


def row_evidence(job):
    rows = list(job.rows.order_by("row_no").values(
        "id", "tenant_id", "row_no", "commit_status", "is_valid", "result_ref")[:MAX_RECEIPT_ROWS + 1])
    if len(rows) > MAX_RECEIPT_ROWS:
        raise ImportReceiptLimit("IMPORT_RECEIPT_ROW_LIMIT")
    refs = {_uuid(x["result_ref"]) for x in rows if x["commit_status"] == "COMMITTED"}
    refs.discard(None)
    staff = set(str(x) for x in HrStaffMaster.objects.filter(
        tenant_id=job.tenant_id, id__in=refs,
        person_id__tenant_id=job.tenant_id).values_list("id", flat=True))
    source_to_row = {f"import:{job.id}:row:{x['row_no']}": x for x in rows}
    relationships = defaultdict(list)
    for item in HrEmploymentRelationship.objects.filter(tenant_id=job.tenant_id,
            staff_id__tenant_id=job.tenant_id, source_business_type="MIGRATION_VERIFIED",
            source_business_id__in=source_to_row).values("id", "staff_id_id", "source_business_id"):
        relationships[item["source_business_id"]].append(item)
    assignments = defaultdict(list)
    for item in HrStaffAssignment.objects.filter(tenant_id=job.tenant_id,
            employment_relationship_id__tenant_id=job.tenant_id,
            employment_relationship_id__staff_id__tenant_id=job.tenant_id,
            source_business_type="MIGRATION_VERIFIED", source_business_id__in=source_to_row,
            assignment_type="PRIMARY").values("id", "employment_relationship_id_id", "source_business_id",
                "organization_id__tenant_id", "position_id__tenant_id", "post_catalog_id__tenant_id",
                "reporting_staff_id__tenant_id"):
        assignments[item["source_business_id"]].append(item)
    audits = Counter((x["business_id"], str(x["staff_id"])) for x in HrStaffAuditEvent.objects.filter(
        tenant_id=job.tenant_id, action="IMPORT_ROW_COMMITTED", business_type="HR03_IMPORT_ROW",
        business_id__in=[f"{job.id}:{x['row_no']}" for x in rows]).values("business_id", "staff_id"))
    referenced = Counter(_uuid(x["result_ref"]) for x in rows if x["commit_status"] == "COMMITTED")
    evidence = []
    for row in rows:
        codes = []
        ref = _uuid(row["result_ref"])
        complete = False
        audit_ok = False
        if row["tenant_id"] != job.tenant_id:
            codes.append("ROW_TENANT_MISMATCH")
        elif row["commit_status"] not in {"PENDING", "COMMITTED", "FAILED"}:
            codes.append("ROW_STATE_INVALID")
        elif row["commit_status"] == "COMMITTED":
            if not row["is_valid"]:
                codes.append("COMMITTED_ROW_INVALID")
            if not ref:
                codes.append("LEGACY_OR_INVALID_RECEIPT")
            elif ref not in staff:
                codes.append("PERSON_OR_STAFF_MISSING")
            else:
                source = f"import:{job.id}:row:{row['row_no']}"
                rels = relationships.get(source, [])
                assigns = assignments.get(source, [])
                if referenced[ref] != 1:
                    codes.append("DUPLICATE_STAFF_RECEIPT")
                if len(rels) != 1 or str(rels[0]["staff_id_id"]) != ref:
                    codes.append("EMPLOYMENT_SOURCE_MISMATCH")
                elif len(assigns) != 1 or assigns[0]["employment_relationship_id_id"] != rels[0]["id"]:
                    codes.append("ASSIGNMENT_SOURCE_MISMATCH")
                if any(item[key] not in (None, job.tenant_id) for item in assigns for key in (
                        "organization_id__tenant_id", "position_id__tenant_id",
                        "post_catalog_id__tenant_id", "reporting_staff_id__tenant_id")):
                    codes.append("ASSIGNMENT_REFERENCE_SCOPE_MISMATCH")
                complete = not codes
                audit_ok = audits[(f"{job.id}:{row['row_no']}", ref)] == 1
                if not audit_ok:
                    codes.append("ROW_AUDIT_MISSING_OR_DUPLICATE")
        elif not row["is_valid"]:
            codes.append("SOURCE_ROW_REQUIRES_CORRECTION")
        else:
            codes.append("ROW_NOT_COMMITTED")
        evidence.append({"rowNo": row["row_no"], "commitStatus": row["commit_status"],
            "valid": bool(row["is_valid"]), "sameSchool": row["tenant_id"] == job.tenant_id,
            "authorityComplete": complete, "auditComplete": audit_ok,
            "verified": complete and audit_ok, "issues": codes})
    return evidence


def inspect_locked_job(job, *, include_issues=True):
    """One authority check used by status, commit result and receipt export.

    Callers lock the job in a transaction, as the existing row writer does.
    This is historical-source evidence, never a current-employment assertion.
    """
    rows = row_evidence(job)
    committed = sum(x["commitStatus"] == "COMMITTED" for x in rows)
    failed = sum(not x["valid"] and x["commitStatus"] != "COMMITTED" for x in rows)
    pending = len(rows) - committed - failed
    verified = sum(x["verified"] for x in rows)
    accounting_issues = []
    if len(rows) != job.total_rows:
        accounting_issues.append("JOB_TOTAL_MISMATCH")
    if job.failed_rows != failed:
        accounting_issues.append("JOB_FAILED_COUNT_MISMATCH")
    if any(not x["sameSchool"] for x in rows):
        accounting_issues.append("ROW_TENANT_MISMATCH")
    if any((x["commitStatus"] == "COMMITTED" and not x["valid"])
           or x["commitStatus"] not in {"PENDING", "COMMITTED", "FAILED"}
           or (x["commitStatus"] == "FAILED" and x["valid"]) for x in rows):
        accounting_issues.append("ROW_STATE_INVALID")
    finished = job.status in {"COMPLETED", "PARTIAL_FAILED"}
    if finished and (pending != 0 or (job.status == "COMPLETED" and failed != 0)
            or (job.status == "PARTIAL_FAILED" and failed == 0)):
        accounting_issues.append("JOB_TERMINAL_STATE_MISMATCH")
    # Historical ledgers without these checkpoint counters are not fabricated.
    if not isinstance(job.checkpoint, dict):
        accounting_issues.append("CHECKPOINT_FORMAT_INVALID")
    checkpoint = job.checkpoint if isinstance(job.checkpoint, dict) else {}
    if finished:
        for name, actual in (("committed_rows", committed), ("failed_rows", failed)):
            if name in checkpoint and (type(checkpoint[name]) is not int or checkpoint[name] != actual):
                accounting_issues.append("CHECKPOINT_COUNT_MISMATCH:" + name)
    balanced = not accounting_issues
    status = "VERIFIED" if finished and balanced and committed > 0 and verified == committed and pending == 0 else "NOT_VERIFIED"
    return {"schemaVersion": "hr03.migration-receipt.1", "jobId": str(job.id),
        "observedAt": datetime.now(timezone.utc).isoformat(), "jobStatus": job.status,
        "verificationStatus": status, "accountingBalanced": balanced, "accountingIssues": accounting_issues,
        "total": job.total_rows, "ledgerRows": len(rows), "committed": committed,
        "failed": failed, "pending": pending, "verified": verified, "rows": rows,
        "legacyUnverifiedRows": sum("LEGACY_OR_INVALID_RECEIPT" in x["issues"] for x in rows),
        "errorRows": safe_issue_rows(job) if include_issues else [],
        "scope": "IMPORTED_ROWS_FOUR_LAYER_AND_ROW_AUDIT",
        "limits": ["错误行未导入，不代表原文件全部成功", "历史来源核对，不证明当前有效任职、账号或生产验收"]}


@transaction.atomic
def migration_receipt(*, tenant_id, job_id):
    job = HrImportJob.objects.select_for_update().get(tenant_id=tenant_id, id=job_id)
    return inspect_locked_job(job)


def receipt_workbook(receipt, issues=()):
    # Runtime uses the application's existing openpyxl dependency. Every text
    # cell is a literal: untrusted labels cannot become Excel formulas.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    summary = [["项目", "结果"], ["导入任务", receipt["jobId"]],
        ["核对状态", receipt["verificationStatus"]], ["任务状态", receipt["jobStatus"]],
        ["源文件行数", receipt["total"]], ["已提交", receipt["committed"]],
        ["错误未写入", receipt["failed"]], ["待处理", receipt["pending"]],
        ["四层事实及审计一致", receipt["verified"]], ["行数守恒", receipt["accountingBalanced"]],
        ["守恒异常", " / ".join(receipt.get("accountingIssues", [])) or "无"],
        ["核对时间", receipt["observedAt"]], ["核对边界", "；".join(receipt["limits"])]]
    detail = [["原文件行号", "写入状态", "四层来源一致", "逐行审计一致", "核对结果", "需处理原因"]]
    detail += [[x["rowNo"], x["commitStatus"], x["authorityComplete"], x["auditComplete"],
        "一致" if x["verified"] else "需核对", " / ".join(ISSUE_LABELS.get(code,code) for code in x["issues"])] for x in receipt["rows"]]
    errors = [["原文件行号", "字段", "错误原因", "错误代码"], *[list(x) for x in issues]]
    for name, matrix in [("迁入核对", summary), ("逐行结果", detail), ("错误行", errors)]:
        sheet = wb.create_sheet(name)
        for values in matrix:
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str): cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="255C91")
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 24 if col < sheet.max_column else 62
        for row in range(2, sheet.max_row + 1): sheet.row_dimensions[row].height = 34
    output=BytesIO(); wb.save(output); return output.getvalue()
