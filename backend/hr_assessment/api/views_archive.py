"""Read, purpose-audited export and single-sample checking of frozen HR12 archives.

R11 evidence projection adapted to existing Django context/permissions. Does not
compute HR scores, change finalized facts, serve HR09 file references or certify
that all procurement samples have passed. See docs/r11_fusion/PROVENANCE.md.
"""
from __future__ import annotations

import io
import json
from decimal import Decimal, InvalidOperation
from django.db import transaction
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_http_methods
from hr_assessment.api.response import api_error, api_success
from hr_assessment.models import HrAssessmentArchivePackage, HrAssessmentArchiveAccessAudit
from hr_assessment.permissions import require_assessment_permission
from hr_assessment.services.archive_evidence import ArchiveEvidenceError, archive_projection, digest

READ_PERMISSIONS = ('hr.assessment.archive_manager', 'hr.assessment.auditor')


def _error(code, message, status):
    response = JsonResponse(api_error(code, message, http_status=status), status=status)
    response['Cache-Control'] = 'private, no-store'
    return response


def _context(request, result_id, version):
    purpose = (request.headers.get('X-HR-Access-Reason') or '').strip()
    if not purpose or len(purpose) > 500 or any(ord(c) < 32 for c in purpose):
        raise ArchiveEvidenceError('ASSESSMENT_ARCHIVE_ACCESS_REASON_REQUIRED', '请填写1至500字的查阅或导出用途。', 400)
    if version < 1 or version > 32767:
        raise ArchiveEvidenceError('ASSESSMENT_ARCHIVE_VERSION_INVALID', '归档版本无效。', 400)
    archive = HrAssessmentArchivePackage.objects.select_related('result').filter(
        tenant_id=request.tenant_id, result_id=result_id, result_version=version,
        archive_status='ARCHIVED',
    ).first()
    if archive is None:
        raise ArchiveEvidenceError('ASSESSMENT_ARCHIVE_NOT_FOUND', '当前授权学校中没有该归档版本。', 404)
    return archive, archive_projection(archive), purpose


def _audit(request, archive, purpose, action, details=None):
    # Every delivery reauthorizes via the original decorator. Never release data
    # on audit persistence failure. Only references/digests, no raw comparisons.
    with transaction.atomic():
        HrAssessmentArchiveAccessAudit.objects.create(
            tenant_id=request.tenant_id, archive=archive, actor_user_id=request.user.pk,
            action=action, purpose=purpose, content_hash=archive.content_hash,
            request_id=str(request.headers.get('X-Request-ID', ''))[:128],
            details_json=details or {},
        )


def _json(data):
    response = JsonResponse(api_success(data=data))
    response['Cache-Control'] = 'private, no-store'
    response['X-Content-Type-Options'] = 'nosniff'
    return response


@require_assessment_permission(READ_PERMISSIONS, sensitive=True)
@require_http_methods(['GET'])
def evidence(request, result_id, version):
    try:
        archive, data, purpose = _context(request, result_id, version)
        _audit(request, archive, purpose, 'VIEW')
        return _json(data)
    except ArchiveEvidenceError as exc:
        return _error(exc.code, str(exc), exc.status)
    except Exception:
        return _error('ASSESSMENT_ARCHIVE_READ_UNAVAILABLE', '归档读取或审计记录暂不可用，未返回档案内容。', 503)


def _text(value):
    """XLSX text, never a formula; do not alter numeric cells."""
    value = str(value if value is not None else '')
    if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
        value = "'" + value
    # Excel XML cannot contain these control characters.
    return ''.join(c for c in value if ord(c) >= 32 or c in '\t\n\r')[:32767]


def build_workbook(data, purpose):
    # The application already depends on openpyxl. Keep the deployed dependency;
    # do not add a separate document engine or send personnel data externally.
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    wb.remove(wb.active)
    def sheet(name, rows, widths):
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append([v if isinstance(v, (Decimal, int)) and not isinstance(v, bool) else _text(v) for v in row])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='245B91')
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for i, width in enumerate(widths, 1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(i)].width = width
        return ws
    result, subject, cycle = data['result'], data.get('subject', {}), data.get('cycle', {})
    sheet('归档说明', [('项目', '归档时的值或说明'), ('归档编号', data['archivePackageId']),
        ('结果版本', data['version']), ('归档时间', data['archivedAt']), ('档案指纹', data['contentHash']),
        ('姓名', subject.get('display_name', '旧归档未冻结，不补写当前值')),
        ('工号', subject.get('staff_code', '')), ('组织', subject.get('org_name', '')),
        ('考核周期', cycle.get('nameAtFinalization', '旧归档未冻结')),
        ('用途', purpose), ('范围', data.get('boundary') or data['identityNote']),
        ('验收限制', '仅此归档的结构化证据，不是全系统或全部101条验收证书。')], [23, 85])
    score = result.get('calculatedScore')
    sheet('结果与规则', [('结果ID', '版本', '等级', '分数', '政策版本', '规则版本', '聚合方式'),
        (data['resultId'], data['version'], result.get('gradeCode'), Decimal(str(score)) if score is not None else '未记录',
         data['calculation'].get('policyVersionId', ''), data['calculation'].get('resultRuleVersionId', ''), data['calculation'].get('scoreAggregation', ''))],
        [40, 10, 16, 14, 40, 40, 22])
    sheet('来源版本', [('来源域', '对象类型', '对象编号', '来源版本', '来源时点', '状态', '快照指纹')] + [
        tuple(x.get(k, '') for k in ('providerType','objectType','objectId','sourceVersion','sourceAsOf','status','snapshotHash')) for x in data['sources']],
        [18, 24, 40, 26, 28, 16, 68])
    sheet('附件版本', [('附件ID', '资质ID', '资质名称', '版本', '摘要', '核对范围')] + [
        tuple(x.get(k, '') for k in ('documentId','credentialId','credentialName','version','sha256','verificationStatus')) for x in data['attachments']],
        [40, 40, 28, 10, 68, 46])
    sheet('纪要与缺口', [('类别', '编号', '说明/原文件名', '文件摘要')] + [
        ('审定纪要', x['id'], x['original_filename'], x['sha256']) for x in data.get('decisionDocuments', [])] + [
        ('资质附件缺口', x.get('credentialId', ''), x['status'], '') for x in data.get('attachmentGaps', [])], [20, 40, 56, 68])
    output = io.BytesIO(); wb.save(output)
    return output.getvalue()


@require_assessment_permission('hr.assessment.archive.export', sensitive=True)
@require_http_methods(['GET'])
def export_evidence(request, result_id, version):
    try:
        archive, data, purpose = _context(request, result_id, version)
        payload = build_workbook(data, purpose)
        _audit(request, archive, purpose, 'EXPORT', {'format': 'xlsx', 'sheets': 5})
        response = HttpResponse(payload, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="hr12-archive-{archive.id}-v{version}.xlsx"'
        response['Cache-Control'] = 'private, no-store'; response['X-Content-Type-Options'] = 'nosniff'
        return response
    except ArchiveEvidenceError as exc:
        return _error(exc.code, str(exc), exc.status)
    except Exception:
        return _error('ASSESSMENT_ARCHIVE_EXPORT_UNAVAILABLE', '导出或审计记录暂不可用，未发送文件。', 503)


@require_assessment_permission(READ_PERMISSIONS, sensitive=True)
@require_http_methods(['POST'])
def compare_evidence(request, result_id, version):
    """Compare one manually supplied benchmark, no score/rule/history mutations."""
    try:
        archive, data, purpose = _context(request, result_id, version)
        if len(request.body) > 4096:
            raise ArchiveEvidenceError('ASSESSMENT_SAMPLE_INVALID', '样本请求过大。', 400)
        body = json.loads(request.body)
        if not isinstance(body, dict) or set(body) - {'expectedScore','expectedGrade','sampleRef'}:
            raise ArchiveEvidenceError('ASSESSMENT_SAMPLE_INVALID', '只接受人工分数、等级和样本依据编号。', 400)
        reference = body.get('sampleRef')
        if not isinstance(reference, str) or not reference.strip() or len(reference) > 200:
            raise ArchiveEvidenceError('ASSESSMENT_SAMPLE_REFERENCE_REQUIRED', '请填写人工核算的依据编号。', 400)
        expected = body.get('expectedScore'); grade = body.get('expectedGrade')
        if isinstance(expected, bool) or expected is None or not isinstance(grade, str) or not grade.strip() or len(grade) > 50:
            raise ArchiveEvidenceError('ASSESSMENT_SAMPLE_INVALID', '人工分数、等级必须明确填写。', 400)
        score = Decimal(str(expected))
        if not score.is_finite() or abs(score) > Decimal('999999999') or score.quantize(Decimal('.01')) != score:
            raise ArchiveEvidenceError('ASSESSMENT_SAMPLE_INVALID', '人工分数必须为至多两位小数的有限数值。', 400)
        actual = data['result'].get('calculatedScore')
        score_match = actual is not None and Decimal(str(actual)) == score
        grade_match = data['result'].get('gradeCode') == grade.strip()
        status = 'NOT_EVALUATED' if actual is None else ('MATCH' if score_match and grade_match else 'DIFFERENT')
        response = {'status': status, 'scoreMatches': score_match, 'gradeMatches': grade_match,
                    'archivedScore': actual, 'archivedGrade': data['result'].get('gradeCode'),
                    'sampleRef': reference.strip(), 'archiveHash': archive.content_hash,
                    'scope': '仅本次人工输入的一份样本；不更改结果，不代表学校签认或全部样本100%一致。'}
        _audit(request, archive, purpose, 'COMPARE', {'sampleHash': digest(body), 'status': status})
        return _json(response)
    except ArchiveEvidenceError as exc:
        return _error(exc.code, str(exc), exc.status)
    except (json.JSONDecodeError, UnicodeDecodeError, InvalidOperation, ValueError, TypeError):
        return _error('ASSESSMENT_SAMPLE_INVALID', '样本内容格式无效。', 400)
    except Exception:
        return _error('ASSESSMENT_ARCHIVE_COMPARE_UNAVAILABLE', '样本核对或审计暂不可用，未改变归档结果。', 503)
